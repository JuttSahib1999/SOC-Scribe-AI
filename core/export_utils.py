import os
import json
from fpdf import FPDF
from docx import Document

def export_json(incident_data, filepath):
    with open(filepath, 'w') as f:
        json.dump(incident_data, f, indent=4)
    return filepath

def export_markdown(incident_data, filepath):
    md_content = f"# {incident_data.get('title', 'Incident Report')}\n\n"
    md_content += f"## Executive Summary\n{incident_data.get('executive_summary', '')}\n\n"
    md_content += f"## Technical Analysis\n{incident_data.get('technical_analysis', '')}\n\n"
    md_content += f"## Response & Remediation\n{incident_data.get('response_actions', '')}\n"
    
    with open(filepath, 'w', encoding='utf-8') as f:
        f.write(md_content)
    return filepath

def export_docx(incident_data, filepath):
    doc = Document()
    doc.add_heading(incident_data.get('title', 'Incident Report'), 0)
    
    doc.add_heading('Executive Summary', level=1)
    doc.add_paragraph(incident_data.get('executive_summary', ''))
    
    doc.add_heading('Technical Analysis', level=1)
    doc.add_paragraph(incident_data.get('technical_analysis', ''))
    
    doc.add_heading('Response & Remediation', level=1)
    doc.add_paragraph(incident_data.get('response_actions', ''))
    
    doc.save(filepath)
    return filepath